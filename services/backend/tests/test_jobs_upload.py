import io
import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from cci.main import app
from cci.db.session import SessionLocal
from cci.db.models.jobs import JobDescription

client = TestClient(app)


def test_jobs_upload_docx():
    # Create an in-memory DOCX
    document = docx.Document()
    document.add_heading("Lead Backend Infrastructure Engineer", level=1)
    document.add_paragraph("Recruitment Rules & Standards:")
    document.add_paragraph("Must have: 5+ years with Python, FastAPI, and PostgreSQL.")
    document.add_paragraph("Preferred: Kubernetes, Docker, and CI/CD pipelines.")
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)

    response = client.post(
        "/api/v1/jobs/upload",
        data={"title": "Lead Backend Infrastructure Engineer", "role": "backend"},
        files={"file": ("jd_infra.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 201
    data = response.json()
    assert "id" in data
    assert data["title"] == "Lead Backend Infrastructure Engineer"
    assert data["canonical_role"] == "backend"
    assert data["file_name"] == "jd_infra.docx"
    assert data["requirements_count"] > 0
    # INVARIANT: Raw document text and contents are NOT exposed in the response
    assert "raw_text" not in data
    assert "text" not in data
    assert "jd_text" not in data

    # Verify data is stored in backend database
    with SessionLocal() as db:
        jd_record = db.get(JobDescription, data["id"])
        assert jd_record is not None
        assert jd_record.file_name == "jd_infra.docx"
        assert "FastAPI" in jd_record.raw_text
        assert "PostgreSQL" in jd_record.raw_text


def test_jobs_upload_pdf():
    # Create an in-memory PDF
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page()
    page.insert_text(
        (50, 72),
        "Principal Machine Learning Engineer\n\nRecruitment Rules:\n"
        "Must have: PyTorch, transformers, and model evaluation.\n"
        "Preferred: vLLM and Qdrant vector database.",
    )
    stream = io.BytesIO()
    pdf_doc.save(stream)
    pdf_doc.close()
    stream.seek(0)

    response = client.post(
        "/api/v1/jobs/upload",
        data={"title": "Principal Machine Learning Engineer", "role": "ml_engineer"},
        files={"file": ("ml_rules.pdf", stream.getvalue(), "application/pdf")},
    )

    assert response.status_code == 201
    data = response.json()
    assert data["title"] == "Principal Machine Learning Engineer"
    assert data["canonical_role"] == "ml_engineer"
    assert data["file_name"] == "ml_rules.pdf"
    assert data["requirements_count"] > 0
    # INVARIANT: Kept in backend, not shown
    assert "raw_text" not in data
    assert "text" not in data


def test_jobs_upload_unsupported_file_type():
    response = client.post(
        "/api/v1/jobs/upload",
        data={"title": "Invalid Job", "role": "backend"},
        files={"file": ("rules.exe", b"invalid executable binary", "application/octet-stream")},
    )
    assert response.status_code == 400
    assert "Only PDF and DOCX files are supported" in response.json()["detail"]
