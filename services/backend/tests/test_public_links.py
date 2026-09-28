import ipaddress
import time

import httpx
import pytest

from cci.live import public_links as links
from cci.security.ssrf import SSRFSecurityError


def test_transport_pins_destination_and_preserves_tls_name(monkeypatch):
    monkeypatch.setattr(links, 'resolve_and_validate_hostname', lambda host: [ipaddress.ip_address('93.184.216.34')])
    transport = links.PublicTransport()
    transport.transport.close()
    def handle(req):
        assert req.url.host == '93.184.216.34'
        assert req.headers['host'] == 'example.com'
        assert req.extensions['sni_hostname'] == 'example.com'
        return httpx.Response(200, text='public')
    transport.transport = httpx.MockTransport(handle)
    with httpx.Client(transport=transport) as client:
        assert client.get('https://example.com/cert').status_code == 200


@pytest.mark.parametrize('url', ['http://localhost/', 'http://127.0.0.1/', 'http://169.254.169.254/',
    'http://[::1]/', 'https://user:pass@example.com/', 'http://example.com:3000/', 'file:///tmp/x'])
def test_private_or_credential_urls_blocked(url):
    result = links.inspect_link(url, time.monotonic() + 5)
    assert result['status'] in {'security_blocked', 'unavailable'}


def test_redirect_to_private_address_never_sent(monkeypatch):
    calls = []
    def resolve(host):
        if host != 'example.com':
            raise SSRFSecurityError('private')
        return [ipaddress.ip_address('93.184.216.34')]
    monkeypatch.setattr(links, 'resolve_and_validate_hostname', resolve)
    transport = links.PublicTransport()
    transport.transport.close()
    def handle(req):
        calls.append(req.url)
        return httpx.Response(302, headers={'location': 'http://169.254.169.254/latest/meta-data'})
    transport.transport = httpx.MockTransport(handle)
    assert links.inspect_link('https://example.com', time.monotonic() + 5, transport)['status'] == 'security_blocked'
    assert len(calls) == 1


def test_mixed_public_private_dns_never_connects(monkeypatch):
    monkeypatch.setattr(links, 'resolve_and_validate_hostname', lambda host: [
        ipaddress.ip_address('93.184.216.34'), ipaddress.ip_address('127.0.0.1')])
    transport = links.PublicTransport()
    transport.transport.close()
    def forbidden(req):
        pytest.fail('A mixed DNS answer reached the network')
    transport.transport = httpx.MockTransport(forbidden)
    assert links.inspect_link('https://example.com', time.monotonic() + 5, transport)['status'] == 'security_blocked'


def test_page_content_is_observation_not_credential_verification():
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'text/html'},
        text='<title>Python Certificate</title><script>SECRET</script><p>Awarded to Example Candidate for completing the Python course.</p>'))
    result = links.inspect_link('https://coursera.org/verify/abc', time.monotonic() + 5, transport)
    assert result['status'] == 'observed'
    assert result['verification'] == 'public_page_observed'
    assert 'SECRET' not in result['excerpt']
    assert result['kind'] == 'credential'


def test_large_pages_gates_and_unscanned_links_are_explicit(monkeypatch):
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'text/html'}, text='x' * (links.MAX_BYTES + 1)))
    assert links.inspect_link('https://example.com', time.monotonic() + 5, transport)['status'] == 'too_large'
    monkeypatch.setattr(links, 'MAX_LINKS', 0)
    assert links.acquire_public_links(['https://example.com'])[0]['status'] == 'not_scanned'


def test_public_digital_certificate_pdf_text_is_inspected():
    import pymupdf
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((50, 50), 'Example Candidate completed the Python Programming Course Certificate.')
        content = document.tobytes()
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'application/pdf'}, content=content))
    result = links.inspect_link('https://issuer.example/certificate.pdf', time.monotonic() + 5, transport)
    assert result['status'] == 'observed'
    assert 'Example Candidate' in result['excerpt']
    assert result['verification'] == 'public_page_observed'


def test_google_drive_gated_link_does_not_assume_project_file():
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'text/html'},
        text='<html><head><title>Google Drive: Sign-in</title></head><body><p>Sign in to continue to Google Drive</p></body></html>'))
    result = links.inspect_link('https://drive.google.com/file/d/123/view', time.monotonic() + 5, transport)
    assert result['status'] == 'access_restricted'
    assert result['inferred_kind'] == 'cloud_storage'
    assert 'cannot verify whether this is a project artifact' in result['detail']


def test_google_drive_certificate_inferred_as_credential():
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'text/html'},
        text='<html><head><title>AWS_Certified_Solutions_Architect.pdf - Google Drive</title></head><body><p>Certificate details and accreditation ID</p></body></html>'))
    result = links.inspect_link('https://drive.google.com/file/d/456/view', time.monotonic() + 5, transport)
    assert result['status'] == 'observed'
    assert result['inferred_kind'] == 'credential'
    assert 'credential or certificate' in result['detail']


def test_google_drive_source_code_inferred_as_project():
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'text/html'},
        text='<html><head><title>distributed_service_source_code.zip - Google Drive</title></head><body><p>Project code repository archive</p></body></html>'))
    result = links.inspect_link('https://drive.google.com/file/d/789/view', time.monotonic() + 5, transport)
    assert result['status'] == 'observed'
    assert result['inferred_kind'] == 'project'
    assert 'project document or source archive' in result['detail']


def test_google_drive_ambiguous_file_not_falsely_claimed():
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'text/html'},
        text='<html><head><title>shared_notes.pdf - Google Drive</title></head><body><p>General document meeting notes</p></body></html>'))
    result = links.inspect_link('https://drive.google.com/file/d/999/view', time.monotonic() + 5, transport)
    assert result['status'] == 'observed'
    assert result['inferred_kind'] == 'cloud_storage'
    assert 'could be a project artifact, credential, personal notes' in result['detail']


def test_fetch_cloud_file_data_zip_archive():
    import io, zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zf:
        zf.writestr('main.py', 'import fastapi\napp = fastapi.FastAPI()\n')
        zf.writestr('README.md', '# Distributed Cache\nPython and Redis implementation\n')
        zf.writestr('requirements.txt', 'fastapi\nuvicorn\nredis\n')
    zip_bytes = buf.getvalue()

    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'application/zip'}, content=zip_bytes))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/test1234/view', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'project'
    assert result['file_count'] == 3
    assert any(f['name'] == 'main.py' for f in result['files'])
    assert 'Python' in result['technologies']
    assert 'Distributed Cache' in result['excerpt']


def test_fetch_cloud_file_data_pdf_document():
    import pymupdf
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((50, 50), 'AWS Certified Solutions Architect Associate Certificate for Candidate.')
        pdf_bytes = document.tobytes()

    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'application/pdf'}, content=pdf_bytes))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/cert1234/view', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'credential'
    assert result['file_count'] == 1
    assert 'AWS' in result['technologies']
    assert 'Candidate' in result['excerpt']


def test_get_cloud_direct_download_url():
    drive_url = 'https://drive.google.com/file/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/view?usp=sharing'
    assert links.get_cloud_direct_download_url(drive_url) == 'https://drive.google.com/uc?export=download&id=1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms'

    docs_url = 'https://docs.google.com/document/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit'
    assert links.get_cloud_direct_download_url(docs_url) == 'https://docs.google.com/document/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/export?format=txt'

    dropbox_url = 'https://www.dropbox.com/s/sample123/project.zip?dl=0'
    assert links.get_cloud_direct_download_url(dropbox_url) == 'https://www.dropbox.com/s/sample123/project.zip?dl=1'


def test_fetch_cloud_file_data_docx_document():
    """DOCX files on cloud storage should be parsed with the real DOCX parser, extracting text and inferring kind."""
    import io
    import docx
    doc = docx.Document()
    doc.add_paragraph('AWS Certified Solutions Architect Associate')
    doc.add_paragraph('This certifies that the candidate has passed the certification exam.')
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    ct = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': ct}, content=docx_bytes))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/docx1234/view', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'credential', f"Expected credential, got {result['inferred_kind']}"
    assert result['file_count'] == 1
    assert result['files'][0]['type'] == 'docx'
    assert 'AWS' in result['technologies']
    assert 'certification' in result['excerpt'].lower()


def test_fetch_cloud_file_data_docx_project_report():
    """A DOCX containing source code keywords should be inferred as a project document."""
    import io
    import docx
    doc = docx.Document()
    doc.add_paragraph('Technical Specification: Distributed Cache Service')
    doc.add_paragraph('Architecture document and source code reference for the codebase.')
    doc.add_paragraph('Built with Python and Docker on Kubernetes.')
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    ct = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': ct}, content=docx_bytes))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/proj5678/view', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'project', f"Expected project, got {result['inferred_kind']}"
    assert 'Python' in result['technologies']
    assert 'Docker' in result['technologies']


def test_inspect_link_docx_extraction():
    """inspect_link should also handle DOCX files, using cloud_file_verified verification."""
    import io
    import docx
    doc = docx.Document()
    doc.add_paragraph('Certificate of Completion: Advanced React and TypeScript Course')
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    ct = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': ct}, content=docx_bytes))
    import time
    result = links.inspect_link('https://example.com/report.docx', time.monotonic() + 10, transport=transport)
    assert result['status'] == 'observed'
    assert result['verification'] == 'cloud_file_verified'
    assert result['inferred_kind'] == 'credential'
    assert result['files'][0]['type'] == 'docx'


def test_fetch_cloud_file_data_image_certificate():
    """Image files with certificate-related names should be inferred as credentials."""
    fake_png = b'\x89PNG\r\n\x1a\n' + b'\x00' * 500
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'image/png'}, content=fake_png))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/img1234/certificate_aws.png', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'credential'
    assert result['files'][0]['type'] == 'image'


def test_fetch_cloud_file_data_image_generic():
    """Image files without certificate keywords should be inferred as cloud_storage."""
    fake_jpg = b'\xff\xd8\xff\xe0' + b'\x00' * 500
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'image/jpeg'}, content=fake_jpg))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/img5678/photo.jpg', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'cloud_storage'


def test_fetch_cloud_file_data_spreadsheet():
    """Spreadsheet files should be detected and labeled appropriately."""
    fake_xlsx = b'PK\x03\x04' + b'\x00' * 500  # Not a real XLSX but content-type drives detection
    ct = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': ct}, content=fake_xlsx))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/sheet1/data.xlsx', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'cloud_storage'
    assert result['files'][0]['type'] == 'spreadsheet'


def test_fetch_cloud_file_data_presentation():
    """Presentation files should be detected and inferred as portfolio."""
    fake_pptx = b'PK\x03\x04' + b'\x00' * 500
    ct = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': ct}, content=fake_pptx))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/pres1/deck.pptx', transport=transport)
    assert result['status'] == 'fetched'
    assert result['inferred_kind'] == 'portfolio'
    assert result['files'][0]['type'] == 'presentation'


def test_inspect_link_image_accepted():
    """inspect_link should accept and handle image content types instead of rejecting them."""
    fake_png = b'\x89PNG\r\n\x1a\n' + b'\x00' * 500
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'image/png'}, content=fake_png))
    import time
    result = links.inspect_link('https://example.com/badge_certificate.png', time.monotonic() + 10, transport=transport)
    assert result['status'] == 'observed'
    assert result['inferred_kind'] == 'credential'
    assert result['files'][0]['type'] == 'image'


def test_docx_not_mistaken_for_zip():
    """A DOCX file (which starts with PK magic bytes) should be detected as DOCX, not ZIP."""
    import io
    import docx
    doc = docx.Document()
    doc.add_paragraph('Resume of John Smith')
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    # Serve as application/octet-stream to test magic-byte fallback detection
    transport = httpx.MockTransport(lambda req: httpx.Response(200, headers={'content-type': 'application/octet-stream'}, content=docx_bytes))
    result = links.fetch_cloud_file_data('https://drive.google.com/file/d/fallback/view', transport=transport)
    assert result['status'] == 'fetched'
    assert result['files'][0]['type'] == 'docx', f"Expected docx type, got {result['files'][0]['type']}"
    assert 'Resume' in result['excerpt'] or 'John Smith' in result['excerpt']
