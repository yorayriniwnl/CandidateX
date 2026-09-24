"""Synthetic golden CV document generators for testing."""

import io
import fitz  # PyMuPDF
import docx


def create_golden_pdf_with_hidden_links() -> bytes:
    """Generates a valid digital PDF with visible text and hidden embedded hyperlink annotations."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)  # Standard Letter size

    # Header
    page.insert_text((50, 60), "Jordan Example", fontsize=18)
    page.insert_text((50, 75), "SYNTHETIC DEMONSTRATION DATA", fontsize=10)
    page.insert_text((50, 90), "Email: jordan@example.test", fontsize=10)
    page.insert_text((50, 105), "Personal Site: https://jordan-example.vercel.app?utm_source=resume&ref=cv", fontsize=10)

    # Section 1: Links (one visible, one hidden behind anchor text)
    page.insert_text((50, 130), "Portfolio & Code:", fontsize=12)
    page.insert_text((50, 150), "Visible GitHub: https://github.com/jordan-example", fontsize=10)

    # Insert anchor text "Distributed Cache Project" with a hidden URI annotation over it
    rect = fitz.Rect(50, 170, 220, 185)
    page.insert_text((50, 180), "Distributed Cache Project", fontsize=10)
    page.insert_link({
        "kind": fitz.LINK_URI,
        "from": rect,
        "uri": "https://github.com/jordan-example/distributed-cache.git",
    })

    # Another hidden link
    rect2 = fitz.Rect(50, 200, 180, 215)
    page.insert_text((50, 210), "LinkedIn Profile", fontsize=10)
    page.insert_link({
        "kind": fitz.LINK_URI,
        "from": rect2,
        "uri": "https://www.linkedin.com/in/jordan-example?ref=resume_pdf",
    })

    # Section 2: Skills
    page.insert_text((50, 240), "Technical Skills", fontsize=12)
    page.insert_text((50, 260), "Python, FastAPI, PostgreSQL, Docker, Kubernetes, PyTorch", fontsize=10)

    # Section 3: Experience
    page.insert_text((50, 290), "Experience", fontsize=12)
    page.insert_text((50, 310), "Senior Software Engineer - CloudScale Inc (2022 - Present)", fontsize=10)
    page.insert_text((50, 325), "Designed high-throughput microservices handling 20k RPS.", fontsize=10)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_golden_docx_with_links() -> bytes:
    """Generates a valid digital DOCX with paragraph text, table, and embedded hyperlinks."""
    doc = docx.Document()
    doc.add_heading("SYNTHETIC DEMONSTRATION DATA", level=3)
    doc.add_heading("Alex Rivera", level=1)
    doc.add_paragraph("Email: alex@example.test")
    doc.add_paragraph("GitHub: https://github.com/alex-rivera")

    # Add paragraph with hyperlink
    p = doc.add_paragraph("Check out my live demo at: ")
    # Add relationship hyperlink
    part = doc.part
    r_id = part.relate_to(
        "https://bob-demo.fly.dev/api/v1?utm_campaign=hiring",
        docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = docx.oxml.shared.OxmlElement("w:hyperlink")
    hyperlink.set(docx.oxml.ns.qn("r:id"), r_id)
    new_run = docx.oxml.shared.OxmlElement("w:r")
    r_text = docx.oxml.shared.OxmlElement("w:t")
    r_text.text = "Bob Live Demo"
    new_run.append(r_text)
    hyperlink.append(new_run)
    p._p.append(hyperlink)

    # Skills table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Go, TypeScript, SQL"
    table.cell(1, 0).text = "Infrastructure"
    table.cell(1, 1).text = "AWS, Terraform, CI/CD"

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
