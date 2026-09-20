import io
import zipfile
import stat

import docx
import httpx
import pytest
from fastapi.testclient import TestClient

from cci.live import acquisition
from cci.live.intake import parse_resume
from cci.live_app import app
from tests.test_live_analysis import intake, transport, archive, SHA

client = TestClient(app)


def test_docx_and_bare_visible_links_are_parsed():
    document = docx.Document()
    document.add_paragraph('Example Candidate')
    document.add_paragraph('github.com/example')
    document.add_paragraph('Skills')
    document.add_paragraph('Python, SQL')
    stream = io.BytesIO()
    document.save(stream)
    result = parse_resume(stream.getvalue(), 'resume.docx')
    assert result.manifest.github_urls == ['https://github.com/example']
    assert result.manifest.claimed_skills == ['Python', 'SQL']


def test_live_backend_exposes_no_candidate_directory_or_lookup():
    assert client.get('/api/v1/candidates').status_code == 404
    assert client.get('/api/v1/dossier/77777777-7777-7777-7777-777777777777').status_code == 404


@pytest.mark.parametrize('identity', ['', 'unrelated'])
def test_unmatched_or_missing_identity_cannot_receive_repository_credit(monkeypatch, identity):
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(transport))
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example/api'], 'github_identity': identity}).json()
    assert data['dossier']['evidence_records']
    assert all(e['confidence'] == 0 for e in data['dossier']['evidence_records'])
    assert data['dossier']['rci'] is None


@pytest.mark.parametrize('code,expected', [(302, 'unavailable'), (404, 'unavailable'), (429, 'rate_limited')])
def test_redirects_not_followed_and_provider_failures_visible(monkeypatch, code, expected):
    visited = []
    def response(req):
        visited.append(str(req.url))
        return httpx.Response(code, headers={'Location': 'http://169.254.169.254/metadata'})
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(response))
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example/api']}).json()
    assert data['sources'][0]['status'] == expected
    assert len(visited) == 1


def test_profile_expansion_uses_only_declared_account_and_bounds_repositories(monkeypatch):
    visited = []
    def response(req):
        visited.append(str(req.url))
        if req.url.path == '/users/example':
            return httpx.Response(200, json={'login': 'example', 'public_repos': 8})
        if req.url.path == '/users/example/repos':
            return httpx.Response(200, json=[{'name': f'api{i}', 'fork': i == 0} for i in range(8)])
        return httpx.Response(404)
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(response))
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example']}).json()
    assert data['sources'][0]['expanded_repositories'] == [f'https://github.com/example/api{i}' for i in range(1, 7)]
    assert len(data['sources'][0]['inventory']) == 8
    assert data['sources'][0]['profile']['login'] == 'example'
    assert data['sources'][0]['status'] == 'observed'
    assert data['sources'][0]['inventory'][0]['inspection_status'] == 'inventory_only'
    assert len(visited) == 8
    assert all('search' not in url for url in visited)


def test_symlinks_are_omitted_without_reading_target(monkeypatch):
    content = io.BytesIO()
    with zipfile.ZipFile(content, 'w') as z:
        link = zipfile.ZipInfo('root/external.py')
        link.create_system = 3
        link.external_attr = (stat.S_IFLNK | 0o777) << 16
        z.writestr(link, '/etc/passwd')
    def response(req):
        if req.url.host == 'codeload.github.com':
            return httpx.Response(200, content=content.getvalue())
        return transport(req)
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(response))
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example/api']}).json()
    assert data['sources'][0]['files_inspected'] == 0
    assert data['dossier']['evidence_records'] == []


def test_private_repository_not_downloaded_even_with_server_token(monkeypatch):
    calls = []
    def response(req):
        calls.append(str(req.url))
        return httpx.Response(200, json={'private': True})
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(response))
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example/api']}).json()
    assert data['sources'][0]['status'] == 'unavailable'
    assert len(calls) == 1


def test_archive_expansion_limit_and_no_candidate_execution(monkeypatch):
    def response(req):
        if '/git/trees/' in req.url.path:
            return httpx.Response(200, json={'truncated': True})
        if req.url.host == 'codeload.github.com':
            return httpx.Response(200, content=archive(content='raise RuntimeError("candidate code must not run")\n@app.get("/")\ndef endpoint():\n    return {}'))
        return transport(req)
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(response))
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example/api'], 'github_identity': 'example'}).json()
    assert data['sources'][0]['status'] == 'observed'
    monkeypatch.setattr(acquisition, 'MAX_EXPANDED_BYTES', 10)
    data = client.post('/api/v1/live/analyze', json={'intake': intake().json(), 'github_urls': ['https://github.com/example/api']}).json()
    assert data['sources'][0]['status'] == 'too_large'


@pytest.mark.parametrize('claim', ['Java', 'Java development'])
def test_live_claims_do_not_treat_unrelated_backend_evidence_as_java_verification(monkeypatch, claim):
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(transport))
    resume = intake().json()
    resume['manifest']['claimed_skills'] = [claim]
    data = client.post('/api/v1/live/analyze', json={'intake': resume,
        'github_urls': ['https://github.com/example/api'], 'github_identity': 'example'}).json()
    assert data['dossier']['claims_corroboration']
    assert all(c['status'] == 'unknown' and not c['grounding_evidence_ids'] for c in data['dossier']['claims_corroboration'])
