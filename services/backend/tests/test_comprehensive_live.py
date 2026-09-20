import io

import docx

from cci.live.intake import parse_resume
from cci.live.report import build_report


def test_docx_identity_tables_sections_and_skill_groups():
    document = docx.Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'AYUSH ROY\nayush@example.com'
    for line in ['Professional Summary', 'Software developer building public applications.',
                 'Technical Skills', 'DevOps & Tools: Docker, CI/CD, Git',
                 'Currently Learning: AWS (S3, Lambda), RAG',
                 'Projects', 'Example Platform — Real Time API\tJan 2025 – Present',
                 'https://github.com/example/platform | Python, FastAPI',
                 'Built an API with automated tests.', 'Education', 'B.Tech Computer Science',
                 'Certifications', 'Python (Basic) — HackerRank Certified',
                 'Achievements', 'Built a hackathon prototype.']:
        document.add_paragraph(line)
    data = io.BytesIO()
    document.save(data)
    result = parse_resume(data.getvalue(), 'resume.docx')
    assert result.manifest.display_name == 'AYUSH ROY'
    assert 'CI/CD' in result.manifest.claimed_skills
    assert 'DevOps & Tools: Docker' not in result.manifest.claimed_skills
    assert 'AWS (S3, Lambda)' in result.manifest.claimed_skills
    assert len(result.manifest.project_claims) == 1
    assert result.resume_review.sections['certifications'] == ['Python (Basic) — HackerRank Certified']
    assert result.resume_review.sections['education'] == ['B.Tech Computer Science']
    assert result.resume_review.learning_skills == ['AWS (S3, Lambda)', 'RAG']


def test_summary_heading_never_becomes_candidate_name():
    document = docx.Document()
    for line in ['Professional Summary', 'Developer of public projects and tools.', 'Skills', 'Python']:
        document.add_paragraph(line)
    data = io.BytesIO()
    document.save(data)
    assert parse_resume(data.getvalue(), 'resume.docx').manifest.display_name == 'Unknown Candidate'


def test_skill_matching_is_exact_and_preserves_attribution_and_learning():
    document = docx.Document()
    for line in ['Example Candidate', 'Skills', 'Java, JavaScript, Python',
                 'Currently Learning: Docker', 'Machine Learning: Scikit-Learn']:
        document.add_paragraph(line)
    data = io.BytesIO()
    document.save(data)
    intake = parse_resume(data.getvalue(), 'resume.docx')
    sources = [{'url': 'https://github.com/example/api', 'status': 'observed', 'ownership_score': 0,
                'repository_review': {'technologies': [{'name': 'JavaScript', 'path': 'app.js',
                'basis': 'source_file_extension', 'url': 'https://github.com/example/api/blob/sha/app.js'}]}}]
    report = build_report(intake, sources)
    skills = {s['skill']: s for s in report['skills']}
    assert skills['Java']['status'] == 'not_observed'
    assert skills['JavaScript']['status'] == 'repository_only'
    assert skills['Docker']['learning']
    assert not skills['Scikit-Learn']['learning']
    sources[0]['ownership_score'] = .5
    assert build_report(intake, sources)['skills'][1]['status'] == 'repository_support'


def test_certificate_matching_does_not_authenticate_or_use_blocked_pages():
    document = docx.Document()
    for line in ['Example Candidate', 'Certifications', 'Python Programming Certificate']:
        document.add_paragraph(line)
    data = io.BytesIO()
    document.save(data)
    intake = parse_resume(data.getvalue(), 'resume.docx')
    source = {'url': 'https://coursera.org/verify/example', 'status': 'observed', 'kind': 'credential',
              'title': 'Python Programming', 'excerpt': 'Example Candidate completed Python Programming.'}
    credential = build_report(intake, [source])['credentials'][0]
    assert credential['status'] == 'possible_public_match'
    assert credential['matching_pages'][0]['candidate_name_present']
    source['status'] = 'access_restricted'
    assert build_report(intake, [source])['credentials'][0]['status'] == 'unverified'


def test_full_service_retains_excluded_links_and_uses_public_pages_without_scoring(monkeypatch):
    from cci.live.contracts import LiveAnalysisRequest
    from cci.live import service
    document = docx.Document()
    for line in ['Example Candidate', 'Skills', 'Python', 'https://example.com/portfolio']:
        document.add_paragraph(line)
    data = io.BytesIO()
    document.save(data)
    intake = parse_resume(data.getvalue(), 'resume.docx')
    monkeypatch.setattr(service, 'acquire_sources', lambda urls, identity: ([], [], []))
    monkeypatch.setattr(service, 'acquire_public_links', lambda urls: [{'url': u, 'status': 'observed',
        'excerpt': 'Python developer with a public portfolio.'} for u in urls])
    result = service.analyze_resume(LiveAnalysisRequest(intake=intake))
    assert result['analysis']['skills'][0]['status'] == 'public_mention_only'
    assert result['dossier'].rci is None
    result = service.analyze_resume(LiveAnalysisRequest(intake=intake, external_urls=[]))
    assert result['sources'][0]['status'] == 'not_selected'


def test_full_service_exposes_requirement_fit_and_critical_gaps(monkeypatch):
    from cci.live.contracts import LiveAnalysisRequest
    from cci.live import service

    document = docx.Document()
    for line in ['Example Candidate', 'Skills', 'Python']:
        document.add_paragraph(line)
    data = io.BytesIO()
    document.save(data)
    intake = parse_resume(data.getvalue(), 'resume.docx')
    monkeypatch.setattr(service, 'acquire_sources', lambda urls, identity: ([], [], []))
    monkeypatch.setattr(service, 'acquire_public_links', lambda urls: [])

    result = service.analyze_resume(LiveAnalysisRequest(
        intake=intake,
        jd_text='Must have Python.\nMust have PostgreSQL.',
    ))

    role_fit = result['dossier'].role_fit
    assert role_fit.mandatory_total == 2
    assert role_fit.mandatory_unknown == 2
    assert len(role_fit.critical_gaps) == 2
    assert result['analysis']['role_fit']['mandatory_unknown'] == 2


def test_large_repository_falls_back_to_commit_pinned_blobs(monkeypatch):
    import base64
    import hashlib
    import httpx
    from cci.live import acquisition
    from tests.test_live_analysis import transport
    content = b'from fastapi import FastAPI\napp = FastAPI()\n@app.get("/items")\ndef items():\n    return []\n'
    blob = hashlib.sha1(f'blob {len(content)}\0'.encode() + content).hexdigest()
    def respond(req):
        if '/git/trees/' in req.url.path:
            return httpx.Response(200, json={'tree': [{'path': 'app.py', 'type': 'blob', 'mode': '100644', 'size': len(content), 'sha': blob}]})
        if '/git/blobs/' in req.url.path:
            return httpx.Response(200, json={'encoding': 'base64', 'content': base64.b64encode(content).decode()})
        return transport(req)
    monkeypatch.setattr(acquisition, 'MAX_ARCHIVE_BYTES', 10)
    monkeypatch.setattr(acquisition, 'HTTP_TRANSPORT', httpx.MockTransport(respond))
    evidence, _, sources = acquisition.acquire_sources(['https://github.com/example/api'], 'example')
    assert sources[0]['status'] == 'observed'
    assert sources[0]['acquisition_method'] == 'bounded_git_blobs'
    assert sources[0]['files_inspected'] == 1
    assert evidence
    assert sources[0]['repository_review']['technologies'][0]['name'] == 'Python'


def test_bare_portfolio_urls_do_not_extract_email_domains():
    from cci.intake.parsers.pdf import URL_REGEX
    matches = [m.group() for m in URL_REGEX.finditer('Portfolio: example.dev | ayushroy.dev@gmail.com | person@sub.example.com | issuer.org/verify/abc | https://issuer.org?credential=123')]
    assert matches == ['example.dev', 'issuer.org/verify/abc', 'https://issuer.org?credential=123']
